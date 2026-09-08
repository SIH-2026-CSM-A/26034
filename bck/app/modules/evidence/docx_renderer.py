from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from docx import Document
from docx.shared import Pt

from .models import OfficerReportModel


def _build_title_section(doc: Document, report: OfficerReportModel):
    doc.add_heading("Compliance Evidence Report", 0)
    p = doc.add_paragraph()
    p.add_run(f"Report ID: {report.report_id}\n").bold = True
    p.add_run(f"Generated: {report.generated_at}\n")
    p.add_run(f"Rule Set: {report.rule_set_version}\n")
    p.add_run(f"Evidence Hash: {report.evidence_hash or 'Not available'}")


def _build_confirmation_section(doc: Document, report: OfficerReportModel):
    doc.add_heading("Officer Confirmation", level=1)
    p = doc.add_paragraph()
    p.add_run(f"Confirmed by: {report.confirmed_by}\n")
    p.add_run(f"Confirmed at: {report.confirmed_at}\n")
    p.add_run(f"Action: {report.officer_action}\n")
    p.add_run(f"Notes: {report.officer_notes or 'None'}")


def _build_verdict_section(doc: Document, report: OfficerReportModel):
    doc.add_heading("Overall Status", level=1)
    v_p = doc.add_paragraph()
    v_run = v_p.add_run(report.overall_verdict)
    v_run.bold = True
    v_run.font.size = Pt(14)


def _build_declarations_section(doc: Document, report: OfficerReportModel):
    doc.add_heading("Extracted Declarations", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Field"
    hdr_cells[1].text = "Value"
    hdr_cells[2].text = "State"
    hdr_cells[3].text = "Provider"
    hdr_cells[4].text = "Conf"
    for d in report.extracted_declarations:
        row_cells = table.add_row().cells
        row_cells[0].text = d.field_name
        row_cells[1].text = d.declared_value or "N/A"
        row_cells[2].text = d.state
        row_cells[3].text = d.ocr_provider
        row_cells[4].text = f"{d.confidence:.2f}" if d.confidence is not None else "N/A"


def _build_rules_section(doc: Document, report: OfficerReportModel):
    doc.add_heading("Rule Evaluation Checklist", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Rule ID"
    hdr_cells[1].text = "Clause"
    hdr_cells[2].text = "Parameter"
    hdr_cells[3].text = "Required"
    hdr_cells[4].text = "Measured"
    hdr_cells[5].text = "Status"
    hdr_cells[6].text = "Notes"
    for r in report.rule_evaluations:
        measured = r.measured_value
        if measured is None and r.state == "INSUFFICIENT_EVIDENCE":
            measured = "Measurement declined"
        elif measured is None:
            measured = "N/A"

        row_cells = table.add_row().cells
        row_cells[0].text = r.rule_id
        row_cells[1].text = r.clause_reference
        row_cells[2].text = r.parameter_name
        row_cells[3].text = r.required_value or "N/A"
        row_cells[4].text = measured
        row_cells[5].text = r.state
        row_cells[6].text = r.notes or ""


def render_to_docx(report: OfficerReportModel, output_stream: BinaryIO | Path) -> bytes:
    """Renders the normalized OfficerReportModel to an editable DOCX."""
    doc = Document()
    _build_title_section(doc, report)
    _build_confirmation_section(doc, report)
    _build_verdict_section(doc, report)
    _build_declarations_section(doc, report)
    _build_rules_section(doc, report)

    doc.add_heading("Source Evidence", level=1)
    doc.add_paragraph(f"Image Path: {report.source_image_path or 'N/A'}")

    out = BytesIO()
    doc.save(out)
    docx_bytes = out.getvalue()
    if isinstance(output_stream, Path):
        output_stream.write_bytes(docx_bytes)
    return docx_bytes
